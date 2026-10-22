"""Tests for the assets style kit. Run with the agent runtime's Python, which has the image and file libraries:

    hermes/.venv/bin/python packages/harness/presets/assets/test_cuttlely_assets.py
"""

import csv
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import cuttlely_assets as kit  # noqa: E402

ROWS = [["Month", "Orders", "Note"], ["Jan", 12, "=SUM(B2:B3)"], ["Feb", 19, "+1 rush"], ["Mar", 15, None]]


class StyleKitTest(unittest.TestCase):
    def setUp(self):
        self.out = Path(tempfile.mkdtemp(prefix="cuttlely-assets-"))

    def test_palette_and_themes_are_hex_colors(self):
        for color in [*kit.PALETTE.values(), *kit.SERIES, kit.TABLE_THEME["header_fill"], kit.TABLE_THEME["band_fill"]]:
            self.assertRegex(color, r"^#[0-9A-F]{6}$")
        self.assertEqual(kit.FONTS["sans"][:3], ["Inter", "Noto Sans", "DejaVu Sans"])

    def test_png_chart(self):
        fig, _ = kit.bar_chart(["Jan", "Feb", "Mar"], [12, 19, 15], title="Orders")
        path = kit.save_png(fig, self.out / "orders.png")
        self.assertTrue(path.read_bytes().startswith(b"\x89PNG\r\n\x1a\n"))
        fig, _ = kit.line_chart(["Jan", "Feb"], {"Orders": [1, 2], "Returns": [0, 1]})
        self.assertTrue(kit.save_png(fig, self.out / "line.png").stat().st_size > 0)

    def test_pdf_report(self):
        fig, _ = kit.bar_chart(["Jan", "Feb"], [1, 2])
        chart = kit.save_png(fig, self.out / "chart.png")
        path = kit.pdf_report(
            self.out / "report.pdf",
            "Q1 <report>",
            [kit.Heading("Orders"), kit.Text("Up & to the right."), kit.Table(ROWS), kit.Image(str(chart), width_cm=10)],
            subtitle="Made-up data",
        )
        data = path.read_bytes()
        self.assertTrue(data.startswith(b"%PDF-"))

    def test_docx_report(self):
        path = kit.docx_report(self.out / "memo.docx", "Q1 memo", [kit.Heading("Orders"), kit.Text("Steady."), kit.Table(ROWS)])
        with zipfile.ZipFile(path) as archive:
            body = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("Q1 memo", body)
        self.assertIn(kit.TABLE_THEME["header_fill"].lstrip("#"), body)

    def test_xlsx_table(self):
        from openpyxl import load_workbook

        path = kit.xlsx_table(self.out / "orders.xlsx", ROWS, sheet="Orders")
        ws = load_workbook(path)["Orders"]
        self.assertEqual(ws["B2"].value, 12)
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["A1"].fill.fgColor.rgb[-6:], kit.TABLE_THEME["header_fill"].lstrip("#"))
        self.assertEqual(ws.freeze_panes, "A2")

    def test_csv_keeps_formula_like_text_as_text(self):
        path = kit.csv_table(self.out / "orders.csv", ROWS)
        with open(path, newline="", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
        self.assertEqual(rows[0], ["Month", "Orders", "Note"])
        self.assertEqual(rows[1][2], "'=SUM(B2:B3)")
        self.assertEqual(rows[2][2], "'+1 rush")
        self.assertEqual(rows[3][2], "")
        self.assertEqual(kit.csv_cell(-3), -3)


if __name__ == "__main__":
    unittest.main()
