"""Smoke test for the image and file libraries in hermes/.venv.

    hermes/.venv/bin/python scripts/hermes-assets/smoke.py [OUT_DIR]

Imports every library from scripts/hermes-assets/requirements.in, renders a PNG, a PDF, a DOCX, an XLSX and a
PPTX from made-up data, and checks each file's signature. Graphviz (`dot`) and the bundled fonts are checked when
the system has them, which the Docker image does. Plotly static export runs when a Chromium is available
(BROWSER_PATH, set in the image); with HERMES_ASSETS_REQUIRE_SYSTEM=1 a missing Chromium fails unless the image was
built with INCLUDE_BROWSER=0. Exits non-zero on the first failure.
"""

import importlib
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

MODULES = [
    "matplotlib",
    "seaborn",
    "plotly",
    "kaleido",
    "numpy",
    "pandas",
    "networkx",
    "graphviz",
    "diagrams",
    "reportlab",
    "docx",
    "openpyxl",
    "xlsxwriter",
    "pptx",
    "PIL",
]

PNG = b"\x89PNG\r\n\x1a\n"
FONT_FAMILIES = ["Inter", "Noto Sans", "DejaVu Sans"]


def check(path, signature):
    data = Path(path).read_bytes()
    if not data.startswith(signature):
        raise SystemExit(f"{path} does not start with {signature!r}")
    print(f"ok  {Path(path).name} ({len(data)} bytes)")


def check_ooxml(path, part):
    with zipfile.ZipFile(path) as archive:
        if part not in archive.namelist():
            raise SystemExit(f"{path} has no {part}")
    check(path, b"PK\x03\x04")


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp(prefix="hermes-assets-"))
    out.mkdir(parents=True, exist_ok=True)

    for name in MODULES:
        module = importlib.import_module(name)
        print(f"ok  import {name} {getattr(module, '__version__', '')}".rstrip())

    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    import pandas as pd
    import seaborn as sns

    frame = pd.DataFrame({"month": ["Jan", "Feb", "Mar", "Apr"], "orders": [12, 19, 15, 24]})
    sns.set_theme(style="whitegrid")
    fig, ax = plt.subplots(figsize=(4, 3), dpi=100)
    sns.barplot(data=frame, x="month", y="orders", ax=ax)
    ax.plot(np.arange(4), frame["orders"].to_numpy(), color="black")
    fig.savefig(out / "chart.png")
    plt.close(fig)
    check(out / "chart.png", PNG)

    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(out / "report.pdf"), pagesize=A4)
    doc.build([Paragraph("Sample report", styles["Title"]), Table([list(frame.columns)] + frame.values.tolist())])
    check(out / "report.pdf", b"%PDF-")

    import docx

    document = docx.Document()
    document.add_heading("Sample memo", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "month"
    table.rows[0].cells[1].text = "orders"
    document.add_picture(str(out / "chart.png"))
    document.save(out / "memo.docx")
    check_ooxml(out / "memo.docx", "word/document.xml")

    import openpyxl
    import xlsxwriter

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        sheet.append(list(row))
    workbook.save(out / "table.xlsx")
    check_ooxml(out / "table.xlsx", "xl/workbook.xml")
    with xlsxwriter.Workbook(str(out / "chart.xlsx")) as book:
        ws = book.add_worksheet()
        ws.write_column("A1", frame["orders"].tolist())
        chart = book.add_chart({"type": "column"})
        chart.add_series({"values": "=Sheet1!$A$1:$A$4"})
        ws.insert_chart("C1", chart)
    check_ooxml(out / "chart.xlsx", "xl/charts/chart1.xml")
    if openpyxl.load_workbook(out / "table.xlsx").active["B2"].value != 12:
        raise SystemExit("table.xlsx did not round-trip")

    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Sample deck"
    deck.save(out / "deck.pptx")
    check_ooxml(out / "deck.pptx", "ppt/presentation.xml")

    import networkx as nx

    graph = nx.path_graph(4)
    if nx.shortest_path_length(graph, 0, 3) != 3:
        raise SystemExit("networkx returned a wrong path length")
    print("ok  networkx")

    if shutil.which("dot"):
        import graphviz

        source = graphviz.Digraph()
        source.edge("ingest", "report")
        path = source.render(str(out / "graph"), format="png", cleanup=True)
        check(path, PNG)
    elif os.environ.get("HERMES_ASSETS_REQUIRE_SYSTEM") == "1":
        raise SystemExit("dot (Graphviz) is not on PATH")
    else:
        print("skip graphviz render: dot is not on PATH")

    if shutil.which("fc-list"):
        families = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True, check=True).stdout
        for family in FONT_FAMILIES:
            if family not in families:
                if os.environ.get("HERMES_ASSETS_REQUIRE_SYSTEM") == "1":
                    raise SystemExit(f"font family {family} is not installed")
                print(f"skip font {family}: not installed")
            else:
                print(f"ok  font {family}")
    elif os.environ.get("HERMES_ASSETS_REQUIRE_SYSTEM") == "1":
        raise SystemExit("fc-list (fontconfig) is not on PATH")

    browser = os.environ.get("BROWSER_PATH")
    if browser and Path(browser).exists():
        import plotly.graph_objects as go

        figure = go.Figure(go.Bar(x=frame["month"], y=frame["orders"]))
        figure.write_image(out / "plotly.png", width=400, height=300)
        check(out / "plotly.png", PNG)
    elif os.environ.get("HERMES_ASSETS_REQUIRE_SYSTEM") == "1" and not Path("/etc/cuttlely/no-browser").exists():
        raise SystemExit("BROWSER_PATH is not set to a Chromium, and this image was not built with INCLUDE_BROWSER=0")
    else:
        print("skip plotly static export: BROWSER_PATH is not set to a Chromium")

    print(f"Hermes asset libraries ready; files in {out}")


if __name__ == "__main__":
    main()
